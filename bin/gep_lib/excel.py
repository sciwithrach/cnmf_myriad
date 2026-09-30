"""Write results as Excel workbooks that are easy to browse: one sheet per GEP."""
from __future__ import annotations

from pathlib import Path
from typing import Dict

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', start_color='EEDEFF')   # UCL pale purple


def _format_sheet(ws, number_formats: Dict[str, str] | None = None) -> None:
    """Bold header, frozen top row, sensible column widths and number formats."""
    ws.freeze_panes = 'A2'
    header = [c.value for c in ws[1]]
    for col_idx, name in enumerate(header, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        width = max([len(str(name))] + [len(str(c.value)) for c in ws[get_column_letter(col_idx)][1:200] if c.value is not None])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(width + 2, 80)
        fmt = (number_formats or {}).get(name)
        if fmt:
            for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                row[0].number_format = fmt


def write_gep_workbook(results: pd.DataFrame, path: str | Path, info: Dict[str, object], n_geps: int) -> None:
    """One sheet per GEP (`GEP_1`, `GEP_2`, ...) of gene set scores, plus an `info` sheet describing the analysis.

    `results` has columns gep, geneset, score, pval. GEPs with no significant gene sets get a sheet with only headers,
    so the workbook always has one sheet for each of the `n_geps` GEPs.
    """
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        pd.DataFrame({'setting': list(info.keys()), 'value': [str(v) for v in info.values()]}).to_excel(
            writer, sheet_name='info', index=False)
        for gep in range(1, n_geps + 1):
            sub = results[results['gep'] == gep].drop(columns='gep')
            sub.to_excel(writer, sheet_name=f'GEP_{gep}', index=False)
        for ws in writer.book.worksheets:
            _format_sheet(ws, {'score': '0.000', 'pval': '0.00E+00'})


def write_table_workbook(table: pd.DataFrame, path: str | Path, sheet_name: str = 'top_genes') -> None:
    """A single-sheet workbook (used for the top genes table: one column per GEP)."""
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        table.to_excel(writer, sheet_name=sheet_name, index=False)
        _format_sheet(writer.book[sheet_name])
